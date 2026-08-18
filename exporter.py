from pathlib import Path
import cv2
import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

def measure(points, pixel_to_mm):
    if len(points) < 2:
        return pd.DataFrame(columns=["ring_no","year_interval","x_start_px","y_start_px","x_end_px","y_end_px","width_px","width_mm","cumulative_distance_mm"])
    width_px = []
    width_mm = []
    cumulative = []
    total = 0.0
    for i in range(1, len(points)):
        d = float(np.linalg.norm(points[i] - points[i-1]))
        mm = d * pixel_to_mm
        total += mm
        width_px.append(d)
        width_mm.append(mm)
        cumulative.append(total)
    return pd.DataFrame({
        "ring_no": np.arange(1, len(width_mm)+1),
        "year_interval": [str(i) + "-" + str(i+1) for i in range(1, len(width_mm)+1)],
        "x_start_px": points[:-1,0],
        "y_start_px": points[:-1,1],
        "x_end_px": points[1:,0],
        "y_end_px": points[1:,1],
        "width_px": width_px,
        "width_mm": width_mm,
        "cumulative_distance_mm": cumulative
    })

def draw_output(img, p1, p2, ring_points, out_path):
    vis = img.copy()
    h, w = vis.shape[:2]
    scale = max(min(w, h) / 1400.0, 0.75)
    point_radius = max(int(5 * scale), 5)
    font_scale = min(max(0.45 * scale, 0.45), 0.8)
    font_thickness = max(int(2 * scale), 2)
    line_width = max(int(3 * scale), 3)

    cv2.line(vis, tuple(p1.astype(int)), tuple(p2.astype(int)), (0,255,0), line_width)
    for i, p in enumerate(ring_points):
        x, y = int(p[0]), int(p[1])
        cv2.circle(vis, (x, y), point_radius + 2, (255, 255, 255), -1)
        cv2.circle(vis, (x, y), point_radius, (0, 0, 255), -1)
        if (i + 1) % 5 == 0 or i == 0 or i == len(ring_points) - 1:
            label = str(i + 1)
            label_x = min(max(x + int(8 * scale), 4), w - 60)
            label_y = min(max(y - int(8 * scale), 18), h - 8)
            cv2.putText(vis, label, (label_x, label_y), cv2.FONT_HERSHEY_SIMPLEX, font_scale, (255, 255, 255), font_thickness + 2, cv2.LINE_AA)
            cv2.putText(vis, label, (label_x, label_y), cv2.FONT_HERSHEY_SIMPLEX, font_scale, (0, 0, 0), font_thickness, cv2.LINE_AA)
    cv2.imwrite(str(out_path), vis)

    labels_path = Path(out_path).with_name(Path(out_path).stem + "_labels.jpg")
    labeled = draw_label_map(img, p1, p2, ring_points)
    cv2.imwrite(str(labels_path), labeled, [int(cv2.IMWRITE_JPEG_QUALITY), 95])

def draw_boundary_output(img, p1, p2, ring_points, out_path, band_width=21):
    vis = img.copy()
    h, w = vis.shape[:2]
    scale = max(min(w, h) / 1400.0, 0.75)
    line_width = max(int(3 * scale), 3)
    font_scale = min(max(0.45 * scale, 0.45), 0.8)
    font_thickness = max(int(2 * scale), 2)

    direction = p2 - p1
    length = float(np.linalg.norm(direction))
    if length < 1e-6:
        direction = np.array([1.0, 0.0], dtype=np.float32)
    else:
        direction = direction / length
    normal = np.array([-direction[1], direction[0]], dtype=np.float32)
    half = max(float(band_width) * 0.85, 12.0)

    cv2.line(vis, tuple(p1.astype(int)), tuple(p2.astype(int)), (0, 180, 0), line_width)
    for i, p in enumerate(ring_points):
        start = p - normal * half
        end = p + normal * half
        x1 = int(np.clip(start[0], 0, w - 1))
        y1 = int(np.clip(start[1], 0, h - 1))
        x2 = int(np.clip(end[0], 0, w - 1))
        y2 = int(np.clip(end[1], 0, h - 1))
        cv2.line(vis, (x1, y1), (x2, y2), (255, 255, 255), line_width + 2, cv2.LINE_AA)
        cv2.line(vis, (x1, y1), (x2, y2), (0, 0, 230), line_width, cv2.LINE_AA)
        if (i + 1) % 5 == 0 or i == 0 or i == len(ring_points) - 1:
            label = str(i + 1)
            x = int(np.clip(p[0] + normal[0] * (half + 8), 4, w - 60))
            y = int(np.clip(p[1] + normal[1] * (half + 8), 18, h - 8))
            cv2.putText(vis, label, (x, y), cv2.FONT_HERSHEY_SIMPLEX, font_scale, (255, 255, 255), font_thickness + 2, cv2.LINE_AA)
            cv2.putText(vis, label, (x, y), cv2.FONT_HERSHEY_SIMPLEX, font_scale, (0, 0, 0), font_thickness, cv2.LINE_AA)
    cv2.imwrite(str(out_path), vis, [int(cv2.IMWRITE_JPEG_QUALITY), 95])

def draw_detection_mode_panel(point_img_path, boundary_img_path, out_path):
    point_img = cv2.imread(str(point_img_path))
    boundary_img = cv2.imread(str(boundary_img_path))
    if point_img is None or boundary_img is None:
        return
    target_h = max(point_img.shape[0], boundary_img.shape[0])

    def resize_to_height(image, height):
        scale = height / float(image.shape[0])
        width = max(int(image.shape[1] * scale), 1)
        return cv2.resize(image, (width, height), interpolation=cv2.INTER_AREA)

    point_view = resize_to_height(point_img, target_h)
    boundary_view = resize_to_height(boundary_img, target_h)
    gap = 28
    title_h = 78
    panel_h = target_h + title_h
    panel_w = point_view.shape[1] + boundary_view.shape[1] + gap
    panel = np.full((panel_h, panel_w, 3), 245, dtype=np.uint8)
    panel[title_h:, :point_view.shape[1]] = point_view
    panel[title_h:, point_view.shape[1] + gap:] = boundary_view
    cv2.rectangle(panel, (point_view.shape[1], title_h), (point_view.shape[1] + gap, panel_h), (245, 245, 245), -1)
    cv2.putText(panel, "Point detection view", (24, 48), cv2.FONT_HERSHEY_SIMPLEX, 1.15, (20, 20, 20), 2, cv2.LINE_AA)
    cv2.putText(panel, "Boundary detection view", (point_view.shape[1] + gap + 24, 48), cv2.FONT_HERSHEY_SIMPLEX, 1.15, (20, 20, 20), 2, cv2.LINE_AA)
    cv2.imwrite(str(out_path), panel, [int(cv2.IMWRITE_JPEG_QUALITY), 95])

def draw_label_map(img, p1, p2, ring_points):
    h, w = img.shape[:2]
    canvas_w = min(max(w, 1900), 3200)
    image_scale = canvas_w / float(w)
    image_h = int(h * image_scale)
    image_view = cv2.resize(img, (canvas_w, image_h), interpolation=cv2.INTER_CUBIC)

    panel_h = 430 if len(ring_points) <= 120 else 560
    canvas = np.full((image_h + panel_h, canvas_w, 3), 245, dtype=np.uint8)
    canvas[:image_h, :] = image_view

    sp1 = (p1 * image_scale).astype(int)
    sp2 = (p2 * image_scale).astype(int)
    point_radius = max(int(5 * image_scale), 6)
    cv2.line(canvas, tuple(sp1), tuple(sp2), (0, 180, 0), max(int(3 * image_scale), 3))

    for i, p in enumerate(ring_points):
        x, y = (p * image_scale).astype(int)
        cv2.circle(canvas, (int(x), int(y)), point_radius + 2, (255, 255, 255), -1)
        cv2.circle(canvas, (int(x), int(y)), point_radius, (0, 0, 230), -1)
        if (i + 1) % 10 == 0 or i == 0 or i == len(ring_points) - 1:
            cv2.putText(canvas, str(i + 1), (int(x) + 10, int(y) - 8), cv2.FONT_HERSHEY_SIMPLEX, 0.75, (255, 255, 255), 4, cv2.LINE_AA)
            cv2.putText(canvas, str(i + 1), (int(x) + 10, int(y) - 8), cv2.FONT_HERSHEY_SIMPLEX, 0.75, (0, 0, 0), 2, cv2.LINE_AA)

    panel_y = image_h
    cv2.rectangle(canvas, (0, panel_y), (canvas_w, image_h + panel_h), (250, 250, 250), -1)
    cv2.line(canvas, (0, panel_y), (canvas_w, panel_y), (50, 50, 50), 2)
    cv2.putText(canvas, "Numbered ring map - labels show detection order / estimated age", (32, panel_y + 42),
                cv2.FONT_HERSHEY_SIMPLEX, 1.0, (25, 25, 25), 2, cv2.LINE_AA)

    if len(ring_points) == 0:
        cv2.putText(canvas, "No detection points.", (32, panel_y + 110), cv2.FONT_HERSHEY_SIMPLEX, 1.0, (0, 0, 180), 2, cv2.LINE_AA)
        return canvas

    margin = 80
    axis_left = margin
    axis_right = canvas_w - margin
    axis_y = panel_y + panel_h // 2
    direction = p2 - p1
    line_len = float(np.linalg.norm(direction))
    if line_len < 1e-6:
        line_len = 1.0
    direction = direction / line_len

    cv2.line(canvas, (axis_left, axis_y), (axis_right, axis_y), (20, 80, 150), 4)
    cv2.putText(canvas, "pith", (axis_left - 58, axis_y + 8), cv2.FONT_HERSHEY_SIMPLEX, 0.8, (20, 80, 150), 2, cv2.LINE_AA)
    cv2.putText(canvas, "bark", (axis_right - 18, axis_y + 42), cv2.FONT_HERSHEY_SIMPLEX, 0.8, (20, 80, 150), 2, cv2.LINE_AA)

    top_rows = [axis_y - 132, axis_y - 92, axis_y - 52]
    bottom_rows = [axis_y + 58, axis_y + 98, axis_y + 138]
    all_rows = top_rows + bottom_rows

    for i, p in enumerate(ring_points):
        dist = float(np.dot(p - p1, direction))
        t = min(max(dist / line_len, 0.0), 1.0)
        tick_x = int(axis_left + t * (axis_right - axis_left))
        tick_top = axis_y - 18
        tick_bottom = axis_y + 18
        cv2.line(canvas, (tick_x, tick_top), (tick_x, tick_bottom), (0, 0, 220), 2)
        cv2.circle(canvas, (tick_x, axis_y), 5, (0, 0, 220), -1)

        row_y = all_rows[i % len(all_rows)]
        label = str(i + 1)
        (tw, th), baseline = cv2.getTextSize(label, cv2.FONT_HERSHEY_SIMPLEX, 0.72, 2)
        label_x = int(tick_x - tw / 2)
        label_x = max(8, min(label_x, canvas_w - tw - 8))
        box_x1 = label_x - 7
        box_y1 = row_y - th - 7
        box_x2 = label_x + tw + 7
        box_y2 = row_y + baseline + 7
        cv2.line(canvas, (tick_x, axis_y), (label_x + tw // 2, row_y - th // 2), (150, 150, 150), 1)
        cv2.rectangle(canvas, (box_x1, box_y1), (box_x2, box_y2), (255, 255, 255), -1)
        cv2.rectangle(canvas, (box_x1, box_y1), (box_x2, box_y2), (0, 0, 220), 1)
        cv2.putText(canvas, label, (label_x, row_y), cv2.FONT_HERSHEY_SIMPLEX, 0.72, (0, 0, 0), 2, cv2.LINE_AA)

    cv2.putText(canvas, "Note: the top image shows point locations; the lower panel is a readable numbered map of the same detections.",
                (32, image_h + panel_h - 28), cv2.FONT_HERSHEY_SIMPLEX, 0.72, (55, 55, 55), 2, cv2.LINE_AA)
    return canvas

def save_all(df, summary, debug_df, out_dir):
    out_dir = Path(out_dir)
    df.to_csv(out_dir / "tree_ring_width_results.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame([summary]).to_csv(out_dir / "summary_report.csv", index=False, encoding="utf-8-sig")
    age_distance = _age_distance_table(df)
    summary_table = _vertical_summary_table(summary)
    with pd.ExcelWriter(out_dir / "tree_ring_results.xlsx", engine="openpyxl") as writer:
        age_distance.to_excel(writer, sheet_name="age_distances", index=False)
        summary_table.to_excel(writer, sheet_name="summary", index=False)
        df.to_excel(writer, sheet_name="raw_ring_widths", index=False)
        debug_df.to_excel(writer, sheet_name="profile_debug", index=False)
        _format_workbook(writer.book)

def _age_distance_table(df):
    if df.empty:
        return pd.DataFrame(columns=[
            "Estimated Age",
            "Age Interval",
            "Distance Between Rings (mm)",
            "Cumulative Distance (mm)",
        ])
    result = pd.DataFrame({
        "Estimated Age": df["ring_no"].astype(int),
        "Age Interval": df["year_interval"],
        "Distance Between Rings (mm)": df["width_mm"].round(4),
        "Cumulative Distance (mm)": df["cumulative_distance_mm"].round(4),
    })
    return result

def _vertical_summary_table(summary):
    labels = {
        "image": "Image",
        "method": "Method",
        "analysis_profile": "Analysis Profile",
        "detection_display_mode": "Detection Display Mode",
        "estimated_age_years": "Estimated Age / Ring Count",
        "detected_boundaries": "Detected Boundary Points",
        "measured_ring_width_count": "Measured Ring Width Count",
        "total_measured_distance_mm": "Total Measured Distance (mm)",
        "mean_width_mm": "Mean Ring Width (mm)",
        "median_width_mm": "Median Ring Width (mm)",
        "std_width_mm": "Standard Deviation (mm)",
        "variation_coefficient": "Variation Coefficient",
        "min_width_mm": "Minimum Ring Width (mm)",
        "max_width_mm": "Maximum Ring Width (mm)",
        "pixel_to_mm": "Pixel to mm",
        "selected_ruler_px": "Selected Ruler (px)",
        "selected_ruler_mm": "Selected Ruler (mm)",
        "min_ring_mm_parameter": "Minimum Ring Width Parameter (mm)",
        "band_width_px": "Band Width (px)",
        "adaptive_block": "Adaptive Block",
        "adaptive_c": "Adaptive C",
    }
    rows = []
    for key, label in labels.items():
        if key in summary:
            rows.append({"Metric": label, "Value": summary.get(key)})
    return pd.DataFrame(rows)

def _format_workbook(workbook):
    header_fill = PatternFill("solid", fgColor="0F766E")
    header_font = Font(color="FFFFFF", bold=True)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"
        for column_cells in sheet.columns:
            max_length = 0
            column_letter = get_column_letter(column_cells[0].column)
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(value))
            sheet.column_dimensions[column_letter].width = min(max(max_length + 3, 14), 42)
